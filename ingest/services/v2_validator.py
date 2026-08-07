import xlrd
import openpyxl
from ingest.models import DescriptiveMetadata, Sheet, People


def _open_workbook(filename: str):
    """Return a normalized workbook wrapper that exposes sheetnames and sheet access."""
    if filename.endswith('.xlsx'):
        wb = openpyxl.load_workbook(filename, read_only=True, data_only=True)
        return _OpenpyxlWrapper(wb)
    else:
        wb = xlrd.open_workbook(filename)
        return _XlrdWrapper(wb)


class _XlrdWrapper:
    def __init__(self, wb):
        self._wb = wb

    @property
    def sheetnames(self):
        return self._wb.sheet_names()

    def sheet(self, name):
        return _XlrdSheetWrapper(self._wb.sheet_by_name(name))

    def close(self):
        pass


class _XlrdSheetWrapper:
    def __init__(self, sheet):
        self._sheet = sheet

    @property
    def max_column(self):
        return self._sheet.ncols

    def iter_rows(self, min_row=1):
        for i in range(min_row - 1, self._sheet.nrows):
            yield self._sheet.row_values(i)


class _OpenpyxlWrapper:
    def __init__(self, wb):
        self._wb = wb

    @property
    def sheetnames(self):
        return self._wb.sheetnames

    def sheet(self, name):
        return _OpenpyxlSheetWrapper(self._wb[name])

    def close(self):
        self._wb.close()


class _OpenpyxlSheetWrapper:
    def __init__(self, sheet):
        self._sheet = sheet

    @property
    def max_column(self):
        return self._sheet.max_column

    def iter_rows(self, min_row=1):
        for row in self._sheet.iter_rows(min_row=min_row, values_only=True):
            yield row


def run_preflight(collection) -> list:
    """Blocking checks run before any Drive lookup.

    Returns a list of error strings. Empty list = all clear.
    """
    errors = []

    try:
        People.objects.get(auth_user_id=collection.user)
    except People.DoesNotExist:
        errors.append(
            f"No People record found for collection owner ({collection.user}). "
            "Cannot attribute upload."
        )
        return errors

    if collection.locked:
        errors.append("Collection is locked and cannot be modified.")

    v1_count = DescriptiveMetadata.objects.filter(collection=collection).count()
    if v1_count == 0:
        errors.append("No V1 metadata (DescriptiveMetadata) found for this collection.")

    v2_count = Sheet.objects.filter(collection=collection).count()
    if v2_count > 0:
        errors.append(
            f"V2 metadata (Sheet) already exists ({v2_count} record(s)). "
            "Remove existing V2 metadata before running this workflow."
        )

    return errors


def validate_spreadsheet(filename: str) -> tuple:
    """Validate a downloaded spreadsheet before upload.

    Returns (errors, warnings). Errors block upload; warnings do not.
    """
    errors = []
    warnings = []

    try:
        workbook = _open_workbook(filename)
    except Exception as e:
        errors.append(f"Could not open spreadsheet: {e}")
        return errors, warnings

    sheet_names = workbook.sheetnames

    if 'README' not in sheet_names:
        errors.append(
            "File appears to be V1 format (no README sheet found). "
            "Upload a V2 spreadsheet."
        )
        return errors, warnings

    if 'SWC' not in sheet_names:
        warnings.append(
            "SWC sheet is missing. If this collection has SWC data, "
            "add the SWC sheet before uploading."
        )

    if 'Dataset' not in sheet_names:
        errors.append("Dataset sheet not found.")
        return errors, warnings

    try:
        ds_sheet = workbook.sheet('Dataset')
        if ds_sheet.max_column <= 15:
            warnings.append(
                "Dataset sheet has 15 or fewer columns — the BILDID column "
                "(column 16) may be missing. Verify before proceeding."
            )
    except Exception:
        pass

    if 'Contributors' not in sheet_names:
        errors.append("Contributors sheet not found.")
        return errors, warnings

    try:
        contrib_sheet = workbook.sheet('Contributors')
        for row in contrib_sheet.iter_rows(min_row=7):
            if len(row) < 6:
                continue
            name_type = row[3] or ''
            name_id = row[4] or ''
            name_id_scheme = row[5] or ''
            if name_type == 'Personal' and (not name_id or not name_id_scheme):
                warnings.append(
                    "Contributors row: nameIdentifier or "
                    "nameIdentifierScheme is empty for a Personal contributor."
                )
    except Exception:
        pass

    workbook.close()
    return errors, warnings
